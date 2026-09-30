from __future__ import annotations

import io
from collections.abc import AsyncIterator
from datetime import datetime

import openpyxl

from app.connectors.base import Connector, SyncResult
from app.services.graph_service import GraphService


class ExcelConnector(Connector):
    """Excel connector using Microsoft Graph API (OneDrive/SharePoint files)."""

    def __init__(self, graph_service: GraphService, account_id: str, user_id: str):
        self.graph_service = graph_service
        self.account_id = account_id
        self.user_id = user_id
        self._workbook_cache: dict[str, openpyxl.Workbook] = {}

    @property
    def name(self) -> str:
        return "excel"

    @property
    def required_scopes(self) -> list[str]:
        return [
            "Files.Read",
            "Files.Read.All",
            "Files.ReadWrite",
            "Files.ReadWrite.All",
            "Sites.Read.All",
            "Sites.ReadWrite.All",
        ]

    async def authenticate(self, credentials: dict) -> bool:
        return True

    def _get_account(self):
        from app.models.account import Account
        account = Account()
        account.id = self.account_id
        account.user_id = self.user_id
        return account

    async def test_connection(self) -> bool:
        try:
            async with self.graph_service.create_client(self._get_account()) as client:
                await client.get_drive()
                return True
        except Exception:
            return False

    async def sync_incremental(
        self,
        since: datetime | None = None,
        cursor: str | None = None
    ) -> AsyncIterator[SyncResult]:
        """Sync Excel files from OneDrive/SharePoint."""
        result = SyncResult()

        async with self.graph_service.create_client(self._get_account()) as client:
            # Search for Excel files
            async for page in client.search_drive("filetype:xlsx OR filetype:xlsm", {"$top": 100}):
                for item in page.get("value", []):
                    result.items_processed += 1
                    result.items_created += 1
                    yield result

    async def get_item(self, item_id: str) -> dict | None:
        """Get Excel file metadata."""
        try:
            async with self.graph_service.create_client(self._get_account()) as client:
                return await client.get(f"/me/drive/items/{item_id}")
        except Exception:
            return None

    async def search(self, query: str, limit: int = 50) -> list[dict]:
        """Search Excel files."""
        results = []
        try:
            async with self.graph_service.create_client(self._get_account()) as client:
                async for page in client.search_drive(query, {"$top": limit}):
                    results.extend(page.get("value", []))
        except Exception:
            pass
        return results[:limit]

    # Excel-specific operations
    async def read_workbook(self, file_id: str, sheet_name: str | None = None) -> dict:
        """Read Excel workbook content."""
        try:
            async with self.graph_service.create_client(self._get_account()) as client:
                # Download file content
                content = await client.get(f"/me/drive/items/{file_id}/content")

                # Parse with openpyxl
                workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

                sheets_data = {}
                target_sheets = [sheet_name] if sheet_name else workbook.sheetnames

                for sheet_name in target_sheets:
                    sheet = workbook[sheet_name]
                    sheets_data[sheet_name] = self._sheet_to_dict(sheet)

                return {
                    "file_id": file_id,
                    "sheets": sheets_data,
                    "sheet_names": workbook.sheetnames,
                }
        except Exception as e:
            return {"error": str(e)}

    async def write_workbook(self, file_id: str, sheets_data: dict) -> dict:
        """Write data to Excel workbook."""
        try:
            async with self.graph_service.create_client(self._get_account()) as client:
                # Download existing or create new
                try:
                    content = await client.get(f"/me/drive/items/{file_id}/content")
                    workbook = openpyxl.load_workbook(io.BytesIO(content))
                except Exception:
                    workbook = openpyxl.Workbook()
                    # Remove default sheet
                    if "Sheet" in workbook.sheetnames:
                        workbook.remove(workbook["Sheet"])

                # Write sheets
                for sheet_name, data in sheets_data.items():
                    if sheet_name in workbook.sheetnames:
                        sheet = workbook[sheet_name]
                    else:
                        sheet = workbook.create_sheet(sheet_name)

                    self._dict_to_sheet(sheet, data)

                # Save to bytes
                output = io.BytesIO()
                workbook.save(output)
                output.seek(0)

                # Upload back
                await client.put(f"/me/drive/items/{file_id}/content", output.read())

                return {"success": True, "file_id": file_id}
        except Exception as e:
            return {"error": str(e)}

    async def import_tasks_from_excel(
        self,
        file_id: str,
        sheet_name: str,
        column_mapping: dict[str, str],
        header_row: int = 1
    ) -> dict:
        """Import tasks from Excel sheet."""
        try:
            workbook_data = await self.read_workbook(file_id, sheet_name)
            if "error" in workbook_data:
                return workbook_data

            sheet_data = workbook_data["sheets"].get(sheet_name, [])
            if not sheet_data:
                return {"error": "Sheet not found or empty"}

            # Skip header row
            rows = sheet_data[header_row:] if header_row > 0 else sheet_data

            tasks = []
            for row in rows:
                task = {}
                for excel_col, task_field in column_mapping.items():
                    # Convert column letter to index
                    col_idx = openpyxl.utils.column_index_from_string(excel_col) - 1
                    if col_idx < len(row):
                        task[task_field] = row[col_idx]

                if task.get("title"):
                    tasks.append(task)

            return {
                "imported": len(tasks),
                "tasks": tasks,
            }
        except Exception as e:
            return {"error": str(e)}

    async def export_tasks_to_excel(
        self,
        file_id: str,
        sheet_name: str,
        tasks: list[dict],
        column_mapping: dict[str, str]
    ) -> dict:
        """Export tasks to Excel sheet."""
        try:
            # Prepare data with headers
            headers = list(column_mapping.keys())
            rows = [headers]

            for task in tasks:
                row = []
                for task_field, excel_col in column_mapping.items():
                    row.append(task.get(task_field, ""))
                rows.append(row)

            sheets_data = {sheet_name: rows}
            return await self.write_workbook(file_id, sheets_data)
        except Exception as e:
            return {"error": str(e)}

    def _sheet_to_dict(self, sheet) -> list[list]:
        """Convert openpyxl sheet to list of lists."""
        data = []
        for row in sheet.iter_rows(values_only=True):
            data.append(list(row))
        return data

    def _dict_to_sheet(self, sheet, data: list[list]) -> None:
        """Write list of lists to openpyxl sheet."""
        for row_idx, row_data in enumerate(data, 1):
            for col_idx, value in enumerate(row_data, 1):
                sheet.cell(row=row_idx, column=col_idx, value=value)

    async def create_workbook(self, name: str, folder_id: str | None = None) -> dict:
        """Create new Excel workbook."""
        try:
            async with self.graph_service.create_client(self._get_account()) as client:
                workbook = openpyxl.Workbook()
                if "Sheet" in workbook.sheetnames:
                    workbook.remove(workbook["Sheet"])

                output = io.BytesIO()
                workbook.save(output)
                output.seek(0)

                # Upload to OneDrive
                path = f"/me/drive/items/{folder_id}:/"
                if not folder_id:
                    path = "/me/drive/root:/"

                result = await client.put(f"{path}{name}.xlsx:/content", output.read())
                return result
        except Exception as e:
            return {"error": str(e)}
